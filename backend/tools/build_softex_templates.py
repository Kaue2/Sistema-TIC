"""Rebuild the seven report templates from the supplied reference PDFs.

Run with the bundled Python runtime, --references <PDF directory> and
--workbook <TRILHA_1.XLS>. Only the report assets/catalog are generated.
"""
import argparse
import io
import json
import re
from pathlib import Path

import pdfplumber
from pypdf import PdfReader
from openpyxl import load_workbook
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
TARGET = ROOT / "backend/src/SistemaTic.Api/Templates/Softex"
RANGES = {"1.13": (6, 8, 15), "1.14": (6, 8, 13), "1.15": (6, 9, 16),
          "2.1": (7, 9, 13), "2.2": (6, 10, 19), "2.4": (6, 11, 19),
          "2.6": (7, 9, 18)}


def clean(text):
    return re.sub(r"\s+", " ", text or "").strip()


def cluster(values):
    groups = []
    for value in sorted(values):
        if not groups or value - groups[-1][-1] > 2:
            groups.append([value])
        else:
            groups[-1].append(value)
    return [sum(group) / len(group) for group in groups]


def extract_rows(pdf, code, labels):
    first, last, expected = RANGES[code]
    rows = []
    for index in range(first - 1, last):
        page = pdf.pages[index]
        xs = cluster(e["x0"] for e in page.edges
                     if e["orientation"] == "v" and e["height"] > 30)
        if len(xs) != 4:
            raise ValueError(f"M{code} page {index + 1}: unexpected columns {xs}")
        ys = cluster(e["top"] for e in page.edges if e["orientation"] == "h"
                     and e["width"] > (xs[2] - xs[1]) * .9
                     and e.get("non_stroking_color") in (0, (0, 0, 0)))
        tables = page.extract_tables({
            "vertical_strategy": "explicit", "explicit_vertical_lines": xs,
            "horizontal_strategy": "explicit", "explicit_horizontal_lines": ys,
        })
        for table in tables:
            for cells in table:
                question, answer, evidence = (clean(c) for c in cells)
                if question in ("Termo:", "Objetivo") or "Trilha de" in question:
                    continue
                # The purpose of the report precedes the per-trail question block.
                if code != "2.1" and question.startswith("Explicação") and "propósito" in question:
                    continue
                starts = any(label.startswith(question) for label in labels if len(question) > 15)
                if starts:
                    rows.append({"label": question, "default": answer})
                elif rows:
                    rows[-1]["label"] = clean(rows[-1]["label"] + " " + question)
                    rows[-1]["default"] = clean(rows[-1]["default"] + " " + answer)
    if len(rows) != expected:
        raise ValueError(f"M{code}: expected {expected} questions, extracted {len(rows)}")
    return rows


def section_text(reader, heading):
    for page in reader.pages:
        lines = [line.strip().rstrip(":") for line in (page.extract_text() or "").splitlines()]
        if heading in lines:
            lines = lines[lines.index(heading) + 1:]
            return "\n".join(line for line in lines if line and not line.isdigit()).strip()
    raise ValueError(f"Missing section: {heading}")


def control(element, tag):
    sdt = OxmlElement("w:sdt")
    props = OxmlElement("w:sdtPr")
    name = OxmlElement("w:tag")
    name.set(qn("w:val"), tag)
    props.append(name)
    sdt.append(props)
    content = OxmlElement("w:sdtContent")
    parent = element.getparent()
    parent.replace(element, sdt)
    content.append(element)
    sdt.append(content)
    return sdt


def tagged_paragraph(doc, tag, text, style=None):
    paragraph = doc.add_paragraph(text, style)
    control(paragraph._p, tag)
    return paragraph


def page_heading(doc, text):
    paragraph = doc.add_paragraph(text, "Heading 1")
    paragraph.paragraph_format.page_break_before = True
    return paragraph


def configure(section, landscape=False):
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width = Cm(29.7 if landscape else 21)
    section.page_height = Cm(21 if landscape else 29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.5)
    section.header_distance = section.footer_distance = Cm(.8)


def fixed_answer(code, number, text):
    # Dates and counts from a sample offer are never institutional defaults.
    if (code, number) == ("2.6", 8):
        return "Data de emissão das declarações: {{emission_date}}."
    if (code, number) == ("2.4", 6):
        return "Datas-chave da coleta registradas: {{collection_dates}}."
    if (code, number) == ("2.4", 15):
        return "Indicadores quantitativos cadastrados: {{indicators}}."
    if (code, number) == ("2.4", 19):
        return ("Os dados coletados serão utilizados para aprimorar as capacitações, "
                "as estratégias de ensino e a experiência educacional.")
    if (code, number) == ("2.6", 17):
        return "As declarações de participação seguem o modelo institucional apresentado nos anexos."
    if (code, number) == ("2.6", 18):
        return "Quantidade de declarações cadastrada para a trilha: {{credential_count}}."
    text = re.sub(r"\b\d{2}/\d{2}/\d{4}\b", "{{emission_date}}", text)
    return text


def build(code, rows, reader, field_map, source, logo):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15
    for name in ("Title", "Heading 1", "Heading 2"):
        doc.styles[name].font.name = "Arial"
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
        doc.styles[name].font.size = Pt(14 if name == "Title" else 12)
    for border in doc.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    summary_style = doc.styles.add_style("Sumario", WD_STYLE_TYPE.PARAGRAPH)
    summary_style.base_style = doc.styles["Heading 1"]
    level = OxmlElement("w:outlineLvl")
    level.set(qn("w:val"), "9")
    summary_style.element.get_or_add_pPr().append(level)
    configure(doc.sections[0])
    header = doc.sections[0].header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.add_run().add_picture(io.BytesIO(logo), width=Cm(1.75))
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    title = doc.add_paragraph("RELATÓRIO DE PRESTAÇÃO DE CONTAS\nSENAC PARA SOFTEX", "Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(95)
    title.paragraph_format.space_after = Pt(20)
    p = doc.add_paragraph("Modalidade: TIC em Trilhas")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("Entrega:")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = tagged_paragraph(doc, "report:stage", "M" + code + " - " + source["name"])
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = tagged_paragraph(doc, "report:date", "{{report_date}}")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(210)
    paragraph = doc.add_paragraph("SUMÁRIO", "Sumario")
    paragraph.paragraph_format.page_break_before = True
    toc = doc.add_paragraph()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    toc.add_run()._r.append(begin)
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = ' TOC \\o "1-2" \\h \\z \\u '
    toc.add_run()._r.append(instruction)
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    toc.add_run()._r.append(separate)
    # Cached section names remain readable in editors that do not update Word fields.
    toc.add_run("RESUMO\nINTRODUÇÃO\nOBJETIVOS\nDESENVOLVIMENTO\nCONCLUSÃO\nREFERÊNCIAS BIBLIOGRÁFICAS\nANEXOS")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    toc.add_run()._r.append(end)
    for heading in ("RESUMO", "INTRODUÇÃO"):
        page_heading(doc, heading)
        doc.add_paragraph(clean(section_text(reader, heading)))
    page_heading(doc, "OBJETIVOS")
    doc.add_paragraph(f"Apresentação das construções que contemplam os critérios da Meta M{code}, "
                      "conforme a Nota Técnica nº 52/2024/GT PPI/SOFTEX.")
    doc.add_paragraph("Este relatório contempla os campos de análise abaixo mencionados:")
    for label in source["all_labels"]:
        doc.add_paragraph(label, "List Number")
    configure(doc.add_section(WD_SECTION_START.NEW_PAGE), True)
    doc.add_paragraph("DESENVOLVIMENTO", "Heading 1")
    # The M2.1 example includes an overview of selected tracks.
    if code == "2.1":
        tagged_paragraph(doc, "report:overview", "{{tracks_overview}}")
    else:
        first = reader.pages[5].extract_text() or ""
        first = first.split("DESENVOLVIMENTO", 1)[1].split("Trilha de", 1)[0]
        doc.add_paragraph(clean(re.sub(r"\n\s*\d+\s*$", "", first)))
    heading = doc.add_paragraph("{{track_title}}", "Heading 2")
    control(heading._p, "track:title")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.autofit = False
    widths = [Cm(7), Cm(12.2), Cm(5.5)]
    for column, width in zip(table.columns, widths):
        column.width = width
    look = table._tbl.tblPr.find(qn("w:tblLook"))
    for attribute in list(look.attrib):
        if attribute != qn("w:val"):
            del look.attrib[attribute]
    for cell, text, width in zip(table.rows[0].cells,
                               ("Pergunta Softex", "Resposta cadastrada", "Evidência"), widths):
        cell.width = width
        cell.text = text
        shade = OxmlElement("w:shd")
        shade.set(qn("w:val"), "clear")
        shade.set(qn("w:fill"), "000000")
        cell._tc.get_or_add_tcPr().append(shade)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for number, row in enumerate(rows, 1):
        question_code = f"M{code}_Q{number:02}"
        field_id = field_map.get(question_code)
        if field_id:
            row["label"] = source["fields"][field_id]
        answer = "" if field_id else fixed_answer(code, number, row["default"])
        cells = table.add_row().cells
        for cell, width in zip(cells, widths):
            cell.width = width
        cells[0].text = row["label"]
        cells[1].text = answer or "-"
        cells[2].text = ""
        control(cells[0].paragraphs[0]._p, "label:" + question_code)
        control(cells[1].paragraphs[0]._p, ("answer:" if field_id else "fixed:") + question_code)
        control(cells[2].paragraphs[0]._p, "evidence:" + question_code)
        for cell in cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1
                for run in p.runs:
                    run.font.size = Pt(9)
    # Wrap the title and its table as one repeatable block.
    block = control(heading._p.getparent().getparent(), "report:trails")
    content = block.find(qn("w:sdtContent"))
    content.append(table._tbl)
    configure(doc.add_section(WD_SECTION_START.NEW_PAGE))
    doc.add_paragraph("CONCLUSÃO", "Heading 1")
    conclusion = clean(section_text(reader, "CONCLUSÃO"))
    # Selection does not imply the lifecycle status of every track.
    conclusion = re.sub(r"As trilhas apresentadas encontram-se finalizada\s*s.*?Cada trilha",
                        "As trilhas selecionadas estão apresentadas em seus respectivos blocos de descrição. Cada trilha",
                        conclusion, flags=re.I)
    doc.add_paragraph(conclusion)
    page_heading(doc, "REFERÊNCIAS BIBLIOGRÁFICAS")
    doc.add_paragraph(section_text(reader, "REFERÊNCIAS BIBLIOGRÁFICAS"))
    page_heading(doc, "ANEXOS")
    tagged_paragraph(doc, "report:annexes", "")
    setting = OxmlElement("w:updateFields")
    setting.set(qn("w:val"), "true")
    doc.settings.element.insert_element_before(setting, "w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr",
                                              "w:compat", "w:docVars", "w:rsids", "m:mathPr")
    doc.core_properties.title = "Relatório de prestação de contas Softex Meta M" + code
    doc.core_properties.author = "Senac"
    doc.save(TARGET / (code + ".docx"))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--references", type=Path, required=True)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--inspect-only", action="store_true")
    args = parser.parse_args()
    mapping_sql = (ROOT / "database/seeds/007_report_remaining_stages.sql").read_text(encoding="utf-8")
    field_map = dict(re.findall(r"\('(M[\d.]+_Q\d+)', '([\d.]+ -\d+)'\)", mapping_sql))
    generated = (ROOT / "frontend/sistema-tic-web/src/data/softexFields.ts").read_text(encoding="utf-8")
    metas = json.loads(generated.split("export const SOFTEX_METAS: SoftexMetaSeed[] = ")[1].strip().rstrip(";"))
    overrides = {item["id"]: item["title"] for meta in metas
                 for item in meta["beforeItems"] + meta["afterItems"]}
    with args.workbook.open("rb") as stream:
        workbook = load_workbook(stream, data_only=True)
    TARGET.mkdir(parents=True, exist_ok=True)
    catalog = []
    inspection = {}
    for code in RANGES:
        sheet = workbook[code]
        labels = {str(sheet.cell(r, 2).value).strip(): clean(sheet.cell(r, 3).value)
                  for r in range(3, sheet.max_row + 1) if sheet.cell(r, 2).value}
        labels.update({key: value for key, value in overrides.items() if key.startswith(code + " -")})
        source = {"fields": labels, "all_labels": list(labels.values()),
                  "name": clean(sheet.cell(1, 1).value).split(" - ", 1)[1]}
        pdf_path = next(p for p in args.references.glob("RELATORIO_DE_PRESTAÇÃO_DE_CONTAS_SENAC_M" + code + "*.pdf")
                        if "(1)" not in p.name)
        reader = PdfReader(pdf_path)
        with pdfplumber.open(pdf_path) as pdf:
            rows = extract_rows(pdf, code, list(labels.values()))
        inspection[code] = [{"number": n, "label": row["label"],
                             "fixed": f"M{code}_Q{n:02}" not in field_map,
                             "default": row["default"] if f"M{code}_Q{n:02}" not in field_map else ""}
                            for n, row in enumerate(rows, 1)]
        if not args.inspect_only:
            logo = reader.pages[0].images[0].data
            build(code, rows, reader, field_map, source, logo)
            catalog += [(f"M{code}_Q{n:02}", row["label"]) for n, row in enumerate(rows, 1)]
    if args.inspect_only:
        print(json.dumps(inspection, ensure_ascii=False, indent=2))
        return
    values = ",\n".join("    ('" + key + "', '" + label.replace("'", "''") + "')" for key, label in catalog)
    seed = ("-- Textos das perguntas dos sete modelos; IDs e vínculos operacionais são preservados.\n"
            "UPDATE report_questions question\n   SET label = source.label\n  FROM (VALUES\n" +
            values + "\n  ) AS source(code, label)\n WHERE question.code = source.code;\n")
    (ROOT / "database/seeds/008_report_question_labels.sql").write_text(seed, encoding="utf-8")
    print("Generated seven DOCX templates and report question labels.")


if __name__ == "__main__":
    main()
